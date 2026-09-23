"""Read-only numeric review: retain original values, inspect grade and gender breakdowns."""
import json, hashlib, re
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[2]
EVIDENCE=ROOT.parent/'docs/planning/incheon-v1/evidence'
OUT=ROOT/'data/interim/incheon/numeric-review'
OUT.mkdir(parents=True,exist_ok=True)
comparison=json.loads((EVIDENCE/'followup/student-reconciliation.json').read_text(encoding='utf-8-sig'))
records=comparison['records']
ice_path=ROOT/'data/raw/incheon/numeric-review/ice-current.xlsx'
old_ice=EVIDENCE/'ice-school-status-20260401.download'
assert ice_path.read_bytes()==old_ice.read_bytes(), 'ICE publication changed: review mapping before proceeding'
kess_path=ROOT/'data/raw/incheon/verified-20260923/kess-2026.download'
with kess_path.open('rb') as f:
    book=load_workbook(f,read_only=True,data_only=True)
    sheet=book['학교별 주요통계']
    rows=sheet.iter_rows(values_only=True)
    header=None; kess={}
    wanted={coordinate_to_tuple(r['kessCell'].split('!')[1])[0] for r in records}
    for i,row in enumerate(rows,1):
        if i==17: header={re.sub(r'\s','',str(v)):j for j,v in enumerate(row) if v is not None}
        if i in wanted: kess[i]=row
        if i>max(wanted):break
    book.close()
ice_book=load_workbook(ice_path,read_only=True,data_only=True)
ice_sheets={n:list(ice_book[n].iter_rows(values_only=True)) for n in ['초등학교','중학교','고등학교','특수학교']}
checks=[]; diffs=[]
for r in records:
    knum,kcol=coordinate_to_tuple(r['kessCell'].split('!')[1]); kn=kess[knum]
    isn,icell=r['iceCell'].split('!'); inum,icol=coordinate_to_tuple(icell)
    ir=ice_sheets[isn][inum-1]
    assert kn[kcol-1]==r['kessStudents'] and ir[icol-1]==r['iceStudents']
    total=kn[header['교원수_총계_계']]; regular=kn[header['교원수_정규_계']]; fixed=kn[header['교원수_기간제교원_계']]
    assert total==regular+fixed
    if r['level']=='특수학교': continue
    grades=6 if r['level']=='초등학교' else 3
    first=22 if grades==6 else 20
    detail=[]
    for g in range(1,grades+1):
        kc=header[f'{g}학년_학생수_계']; kf=header[f'{g}학년_학생수_여']; ic=first-1+(g-1)*2
        detail.append({'grade':g,'kess':kn[kc],'ice':ir[ic],'difference':kn[kc]-ir[ic] if isinstance(kn[kc],(int,float)) and isinstance(ir[ic],(int,float)) else None,
          'kessFemale':kn[kf],'iceFemale':ir[ic+1],
          'kessCell':f'학교별 주요통계!{get_column_letter(kc+1)}{knum}', 'iceCell':f'{isn}!{get_column_letter(ic+1)}{inum}'})
    ksum=sum(d['kess'] for d in detail if isinstance(d['kess'],(int,float))); isum=sum(d['ice'] for d in detail if isinstance(d['ice'],(int,float)))
    checks.append({'id':r['schoolId'],'kessGradeSumMatches':ksum==r['kessStudents'],'iceGradeSumMatches':isum==r['iceStudents']})
    if r['differenceKessMinusIce']:
        diffs.append({**r,'grades':detail,'gradeDifferencesSum':sum(d['difference'] for d in detail) if all(d['difference'] is not None for d in detail) else None,
          'causeStatus':'unconfirmed','finding':'학년별 관측값 차이; 행정적 원인 미확정'})
ice_book.close()
report={'status':'PARTIAL','checkedRecords':len(records),'studentDifferences':len(diffs),'teacherComponentPassed':len(records),
 'gradeSumChecks':{'records':len(checks),'kessPassed':sum(c['kessGradeSumMatches'] for c in checks),'icePassed':sum(c['iceGradeSumMatches'] for c in checks)},
 'sourceHashes':{'kess':hashlib.sha256(kess_path.read_bytes()).hexdigest(),'ice':hashlib.sha256(ice_path.read_bytes()).hexdigest()},
 'iceCurrentDownloadUnchanged':True,'independentTeacherAggregate':'not_acquired','differenceCausesConfirmed':0,
 'differenceDistribution':dict(Counter(abs(r['differenceKessMinusIce']) for r in diffs)),
 'largestDifferences':sorted(diffs,key=lambda r:abs(r['differenceKessMinusIce']),reverse=True)[:3],
 'records':diffs,'failedGradeChecks':[c for c in checks if not(c['kessGradeSumMatches'] and c['iceGradeSumMatches'])],
 'gradeSumPolicy':'Sum populated numeric grade cells; blanks retained as null, not assigned zero. This is an internal sum check, not independent cause verification.'}
historical_controls=[]
for ident,level in [(895,'elem'),(896,'mid'),(897,'high')]:
    source=ROOT/f'data/raw/incheon/numeric-review/ice-statistics-{ident}.html'
    html=source.read_text(encoding='utf-8')
    table=next(t for t in re.findall(r'<table[\s\S]*?</table>',html) if re.search(r'<caption>[^<]*교원수',t))
    years=[int(y) for y in re.findall(r'<th[^>]*>\s*(\d{4})년',table)]
    values=[int(v.replace(',','')) for v in re.findall(r'<td[^>]*>\s*([\d,]+)명',table)]
    assert len(years)==len(values)==10
    for year,value in zip(years,values):
        if year<2022:continue
        annual=json.loads((ROOT/f'data/interim/incheon/history-review/{year}.json').read_text(encoding='utf-8'))
        total=sum(r['values']['teachers'] for r in annual if r['level']==level)
        historical_controls.append({'level':level,'year':year,'kess':total,'icePublished':value,'difference':total-value,
           'url':f'https://www.ice.go.kr/ice/cm/cntnts/cntntsView.do?cntntsId={ident}&mi=11853',
           'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),
           'scope':'separate official publication; not a separate independent survey or 2026 control'})
report['historicalTeacherControls']=historical_controls
(OUT/'numeric-review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['records','sourceHashes','largestDifferences']},ensure_ascii=False))
